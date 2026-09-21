# -*- coding: utf-8 -*-
"""专利增量更新工作流：Google Patents 导出 PDF → about.xlsx

用法（在任意目录执行）：
    python3 update_patents_from_pdf.py \
        --chn /Users/xz/XZhen/Website/raw_data/patents_chn.pdf \
        --eng /Users/xz/XZhen/Website/raw_data/patents_eng.pdf \
        --xlsx /Users/xz/XZhen/Website/raw_data/about.xlsx
    # 预览不写盘：
    python3 update_patents_from_pdf.py --dry-run

输入：
  --chn  Google Patents 检索页导出 PDF（发明人 郑羲光, country:CN, GRANT, CHINESE）
  --eng  Google Patents 检索页导出 PDF（inventor:(Xiguang ZHENG), GRANT, ENGLISH）

更新规则：
  中国专利 sheet
    - 以中文 PDF 为准做增量：新专利号 → 追加；已存在 → 授权日期以 PDF 为准
    - 排除杜比实验室特许公司（与英文清单同族，不录入）
    - 全表按授权日期倒序重排，编号 #N→#1 重编
    - 新专利发明人暂记"郑羲光 等"，备注标注待补全名单
    - 维护性修复：标题粘连发明人名（如"张晨音频处理方法…"）自动移回发明人列
  国际专利 sheet（美国专利）
    - 以英文 PDF 为准做增量（US/EP 均收录，国家列标注）
    - 同名同族只保留一件（先收录者优先），全表连续重编号
    - 中文名称未知的新专利标"—（待翻译）"，请在 xlsx 中补全后重新运行即可生效去重
  说明 sheet 的版本行与各 sheet 计数同步更新。
"""
import argparse
import os
import re
import sys
import unicodedata

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side

# 常见 CJK 异体字（Google Patents PDF 使用兼容部首形式）→ 规范字形
VARIANT_MAP = {
    "⻬": "齐", "⾳": "音", "⽅": "方", "⽐": "比", "⽚": "片",
    "⾼": "高", "⽴": "立", "⽣": "生", "⽤": "用", "⽿": "耳",
    "⻨": "麦", "⻛": "风", "⻓": "长", "⻔": "门", "⻋": "车",
    "⻩": "黄", "⽆": "无", "⾃": "自", "⼀": "一", "⽅": "方",
}

# 已知国际专利中文译名（新增未知专利时先标"—（待翻译）"，补全后重跑本脚本）
KNOWN_CN_NAMES = {
    "US10594283B2": "音频信号响度控制",
    "US10096329B2": "提升音频信号中语音内容的可懂度",
    "US10038961B2": "电声换能器频率响应特性的建模",
    "US9911404B2": "耳机中主动降噪与噪声补偿相结合",
    "US10356524B2": "用户体验导向的音频信号处理",
    "US10750306B2": "用于耳机虚拟化的混响生成",
    "US10848873B2": "方向感知的环绕声回放",
    "US11636836B2": "音频处理方法及电子设备",
    "US11902762B2": "方向感知的环绕声回放",
    "US12335717B2": "使用方向性房间脉冲响应的空间音频重现方法与装置",
    "US10149082B2": "用于耳机虚拟化的混响生成",
    "US9877108B2": "用户体验导向的音频信号处理",
    "US10242659B2": "耳机中主动降噪与噪声补偿相结合",
    "EP3149970B1": "音频信号响度控制",
    "EP3149730B1": "提升音频信号中语音内容的可懂度",
}

RAW_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_XLSX = os.path.join(RAW_DIR, "about.xlsx")
DEFAULT_CHN = os.path.join(RAW_DIR, "patents_chn.pdf")
DEFAULT_ENG = os.path.join(RAW_DIR, "patents_eng.pdf")

CN_SHEET = "中国专利"
INTL_SHEET = "美国专利"  # sheet 名保留历史名称，标题行写作"国际发明专利"
DOLBY_MARK = "实验室特许"  # "杜比实验室特许公司"，避开异体字问题

THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def clean(s):
    """NFKC 归一化 + 常见异体字替换。"""
    if s is None:
        return None
    s = unicodedata.normalize("NFKC", str(s))
    for a, b in VARIANT_MAP.items():
        s = s.replace(a, b)
    return s.strip()


def norm_no(s):
    """专利号归一化：去横线空格、大写，用于跨格式匹配。"""
    return re.sub(r"[-\s]", "", str(s)).upper()


def fmt_cn_no(no):
    """CN115410586B → CN-115410586-B（与 xlsx 既有格式一致）。"""
    m = re.match(r"(CN)(\d+)([A-Z]\d?)$", no)
    return f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else no


def extract_pdf_text(path):
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        return "\n".join(p.extract_text() or "" for p in pdf.pages)


def parse_cn_pdf(path):
    """解析中文 PDF → [{no,title,assignee,y,m,d}]（含杜比，调用方过滤）。"""
    text = extract_pdf_text(path)
    pat = re.compile(
        r"\n([^\n]+)\n(?:Download[^\n]*\n)*"
        r"(?:WO EP US CN JP DK ES HU PL|WO EP US CN JP|WO EP US CN|"
        r"EP US CN|WO CN|WO EP US|CN) • (CN[0-9]+B) • 郑羲光 • ([^\n]+)\n"
        r".*?Granted (\d{4})-(\d{2})-(\d{2})", re.S)
    out = []
    for m in pat.finditer(text):
        title, no, assignee, y, mo, d = m.groups()
        if title.startswith("Download"):  # 版式错位防护
            continue
        out.append({"no": no, "title": clean(title), "assignee": clean(assignee),
                    "y": int(y), "m": int(mo), "d": int(d)})
    return out


def parse_eng_pdf(path):
    """解析英文 PDF → [{no,title,assignee,y,m,d,country}]。"""
    text = extract_pdf_text(path)
    pat = re.compile(
        r"\n([^\n]+)\n(?:Download[^\n]*\n)*"
        r"(?:WO EP US CN JP DK ES HU PL|WO EP US CN JP|WO EP US CN|"
        r"US KR|WO CN|WO EP US|EP US CN|US|EP) • ([A-Z]{2}[0-9]+[AB]\d?) • "
        r"Xiguang ZHENG • ([^\n]+)\n.*?Granted (\d{4})-(\d{2})-(\d{2})", re.S)
    out = []
    for m in pat.finditer(text):
        title, no, assignee, y, mo, d = m.groups()
        if title.startswith("Download"):
            continue
        country = "美国" if no.startswith("US") else ("欧洲" if no.startswith("EP") else no[:2])
        out.append({"no": no, "title": clean(title), "assignee": clean(assignee),
                    "y": int(y), "m": int(mo), "d": int(d), "country": country})
    return out


def style_cell(cell, bold=False):
    cell.border = BORDER
    cell.alignment = WRAP
    if bold:
        cell.font = Font(bold=True)


def rewrite_sheet(ws, header_cols, rows):
    """清空数据区（第 3 行起）按 rows 重写，并物理删除多余旧行。"""
    if ws.max_row > 2:
        ws.delete_rows(3, ws.max_row - 2)
    for r in range(3, ws.max_row + 1):
        for c in range(1, header_cols + 1):
            ws.cell(r, c).value = None
    for i, row in enumerate(rows):
        for c, v in enumerate(row, 1):
            cell = ws.cell(3 + i, c, v)
            style_cell(cell, bold=(c == 1))


def update_cn_sheet(wb, pdf_entries, today, dry_run):
    ws = wb[CN_SHEET]
    pdf = {}
    dolby = 0
    for e in pdf_entries:
        if DOLBY_MARK in e["assignee"]:
            dolby += 1
            continue
        pdf[norm_no(e["no"])] = e
    print(f"[CN] PDF 有效 {len(pdf)} 件（排除杜比 {dolby} 件）")

    records = {}
    for r in range(3, ws.max_row + 1):
        no = ws.cell(r, 2).value
        if not no:
            continue
        k = norm_no(no)
        if k in records:  # 重复行：保留信息更全的一条
            old, new = records[k], [ws.cell(r, c).value for c in range(1, 11)]
            if sum(v is not None for v in new) > sum(v is not None for v in old):
                records[k] = new
            print(f"[CN] 去重: {no}")
            continue
        records[k] = [ws.cell(r, c).value for c in range(1, 11)]

    added, date_fixed = [], []
    for k, e in pdf.items():
        if k in records:
            row = records[k]
            if (row[4], row[5], row[6]) != (e["y"], e["m"], e["d"]):
                date_fixed.append((e["no"], f"{row[4]}-{row[5]}-{row[6]}", f"{e['y']}-{e['m']}-{e['d']}"))
                row[4], row[5], row[6] = e["y"], e["m"], e["d"]
            t = clean(row[2])
            m = re.match(r"^(张晨|张旭|郭亮|郑羲光)[一-鿿]", t)  # 标题粘连发明人名修复
            if m and m.group(1) not in str(row[3]):
                row[2] = t[len(m.group(1)):]
                row[3] = f"{m.group(1)}, {row[3]}"
                print(f"[CN] 修复标题粘连: {e['no']}")
            if e["title"] and e["title"] != t and e["title"] != row[2]:
                row[2] = e["title"]
        else:
            records[k] = [None, fmt_cn_no(e["no"]), e["title"], "郑羲光 等",
                          e["y"], e["m"], e["d"], "发明", "已授权",
                          f"{today} 依 Google Patents 清单新增，发明人全名单待补"]
            added.append(e["no"])

    gone = sorted(set(records) - set(pdf))
    if gone:
        print(f"[CN] 警告: xlsx 有而 PDF 无 {len(gone)} 件，保留不动: {gone}")

    rows = sorted(records.values(), key=lambda x: (x[4], x[5], x[6]), reverse=True)
    n = len(rows)
    for i, row in enumerate(rows):
        row[0] = n - i
    print(f"[CN] 共 {n} 件，新增 {len(added)}: {added}")
    for no, old, new in date_fixed:
        print(f"[CN] 日期修正 {no}: {old} -> {new}")

    if not dry_run:
        rewrite_sheet(ws, 10, rows)
        ws.cell(1, 1).value = f"中国发明专利（已授权）CN Patents Granted（共 {n} 件，不含杜比同族）"
    return n, added


def update_intl_sheet(wb, pdf_entries, today, dry_run):
    ws = wb[INTL_SHEET]
    pdf = {norm_no(e["no"]): e for e in pdf_entries}
    print(f"[INTL] PDF 有效 {len(pdf)} 件")

    rows = []
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 2).value:
            rows.append([ws.cell(r, c).value for c in range(1, 10)])

    added = []
    for k, e in pdf.items():
        if k in {norm_no(row[1]) for row in rows}:
            # 已存在：授权年份以 PDF 为准
            for row in rows:
                if norm_no(row[1]) == k and row[5] != e["y"]:
                    print(f"[INTL] 年份修正 {e['no']}: {row[5]} -> {e['y']}")
                    row[5] = e["y"]
            continue
        cn_name = KNOWN_CN_NAMES.get(e["no"], "—（待翻译）")
        rows.append([None, e["no"], e["title"], cn_name, "Xiguang ZHENG 等",
                     e["y"], e["country"], "已授权",
                     f"{today} 依 Google Patents 新增，发明人全名单待补"])
        added.append(e["no"])

    # 同名同族去重：先收录者优先
    seen, keep, dropped = set(), [], []
    for row in rows:
        key = str(row[3]).strip()
        if key.startswith("—"):  # 未翻译的以英文名称去重
            key = "EN:" + str(row[2]).strip().lower()
        if key in seen:
            dropped.append(row)
        else:
            seen.add(key)
            keep.append(row)
    for d in dropped:
        print(f"[INTL] 同名去重移除: {d[1]} {d[3]}")

    for i, row in enumerate(keep):
        row[0] = i + 1
    us = sum(1 for k in keep if k[6] == "美国")
    ep = sum(1 for k in keep if k[6] == "欧洲")
    print(f"[INTL] 共 {len(keep)} 件（美国 {us} + 欧洲 {ep}），新增 {len(added)}: {added}")

    if not dry_run:
        rewrite_sheet(ws, 9, keep)
        ws.cell(1, 1).value = f"国际发明专利（已授权）International Patents Granted（美国 {us} 件 + 欧洲 {ep} 件）"
    return len(keep), us, ep, added


def main():
    ap = argparse.ArgumentParser(description="Google Patents PDF → about.xlsx 专利增量更新")
    ap.add_argument("--chn", default=DEFAULT_CHN, help="中文专利 PDF 路径")
    ap.add_argument("--eng", default=DEFAULT_ENG, help="英文专利 PDF 路径")
    ap.add_argument("--xlsx", default=DEFAULT_XLSX, help="about.xlsx 路径")
    ap.add_argument("--dry-run", action="store_true", help="只打印变更，不写盘")
    args = ap.parse_args()

    today = __import__("datetime").date.today().strftime("%Y-%m")

    missing = [p for p in (args.chn, args.eng) if not os.path.exists(p)]
    if missing:
        print("找不到输入文件：", missing)
        print("请先从 Google Patents 导出 PDF（检索 inventor:郑羲光 / (Xiguang ZHENG), status:GRANT）放到 raw_data/ 下")
        sys.exit(1)

    wb = load_workbook(args.xlsx)
    cn_n, cn_added = update_cn_sheet(wb, parse_cn_pdf(args.chn), today, args.dry_run)
    intl_n, us, ep, intl_added = update_intl_sheet(wb, parse_eng_pdf(args.eng), today, args.dry_run)

    if not args.dry_run:
        ws = wb["说明"]
        ws.cell(2, 1).value = (
            f"版本专利更新 · {today} · 依 Google Patents 导出 PDF 增量更新"
            f"（中国 {cn_n} 件 / 国际 {intl_n} 件）；杜比 CN 同族不录入中国专利 sheet；"
            f"同名同族国际专利只保留一件")
        ws.cell(13, 1).value = f"  国际专利：编号|专利号|英文名称|中文名称|发明人|授权年份|国家|法律状态|备注（美国 {us} 件 + 欧洲 {ep} 件，同名同族只保留一件）"
        ws.cell(14, 1).value = f"  中国专利：编号|专利号|名称|发明人|授权年|授权月|授权日|专利类型|法律状态|备注（共 {cn_n} 件，编号最新 → #1 最早）"
        wb.save(args.xlsx)
        print("saved:", args.xlsx)
    else:
        print("dry-run 完成，未写盘。")


if __name__ == "__main__":
    main()
